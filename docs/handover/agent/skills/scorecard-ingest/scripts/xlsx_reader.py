"""Read sheet tabs from Excel exports (workaround when G-Sheet API is blocked)."""

from __future__ import annotations

from pathlib import Path


class XlsxError(RuntimeError):
    pass


def read_tab(path: str | Path, tab: str | None = None) -> tuple[str, list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise XlsxError("Install openpyxl to use --mode xlsx") from exc

    book_path = Path(path)
    if not book_path.exists():
        raise XlsxError(f"file not found: {book_path}")

    wb = load_workbook(book_path, read_only=True, data_only=True)
    names = wb.sheetnames
    if not names:
        raise XlsxError(f"no sheets in {book_path}")

    ws = None
    chosen = tab or ""
    if tab:
        if tab in names:
            ws = wb[tab]
            chosen = tab
        else:
            want = tab.strip().lower()
            for name in names:
                if name.strip().lower() == want:
                    ws = wb[name]
                    chosen = name
                    break
    if ws is None:
        ws = wb[names[0]]
        chosen = names[0]

    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell) for cell in row])
    wb.close()
    return chosen, rows
